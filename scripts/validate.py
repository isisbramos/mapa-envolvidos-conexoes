#!/usr/bin/env python3
"""Confere data/mapa.xlsx contra as regras editoriais antes do build.

Uso:
    python3 scripts/validate.py

Sai com código 1 se houver ERRO. AVISO não derruba o build — é a fila do que
precisa de decisão humana. As regras estão escritas em EDITORIAL.md.
"""
import re
import sys
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta a biblioteca openpyxl. Instale com: pip3 install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "data" / "mapa.xlsx"
PHOTOS_DIR = ROOT / "photos"

STATUS_VALIDOS = {"Reportado", "Alegado ou contestado"}
# Estágio em que a ausência de fonte é decisão editorial, não falha:
# pessoas privadas, não nomeadas ou de identidade não confirmada.
ESTAGIOS_SEM_PESQUISA = {"Não pesquisado"}
MIN_FONTES_PERFIL = 2

# Regra F7 do EDITORIAL.md: ligação "Alegado ou contestado" precisa de duas
# fontes registradas nela própria; "Reportado" se apoia nas fontes do perfil.
STATUS_EXIGE_FONTE_PROPRIA = "Alegado ou contestado"
MIN_FONTES_LIGACAO = 2

# Enquanto o preenchimento das ligações não termina, a regra acima é aviso.
# Vire para True quando toda ligação "Alegado ou contestado" tiver duas
# fontes: a partir daí ela derruba o CI e a regra passa a se defender sozinha.
EXIGIR_FONTE_NA_LIGACAO = False

erros = []
avisos = []


def erro(codigo, alvo, msg):
    erros.append(f"ERRO  {codigo}  {alvo}: {msg}")


def aviso(codigo, alvo, msg):
    avisos.append(f"AVISO {codigo}  {alvo}: {msg}")


def txt(v):
    return "" if v is None else str(v).strip()


def conta_fontes(v):
    """Conta URLs distintas no formato 'Título URL | Título2 URL2'."""
    return len(set(re.findall(r"https?://\S+", txt(v))))


def linhas(ws):
    """Linhas com dado, cada uma com o número real dela na planilha.

    A aba Edges tem linhas em branco no meio, então contar só as preenchidas
    daria um número deslocado — e mandar quem revisa para a linha errada é
    pior do que não indicar linha nenhuma.
    """
    return [(i, r) for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2)
            if r and r[0]]


def main():
    if not XLSX_PATH.exists():
        sys.exit(f"Não achei {XLSX_PATH}")
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

    nodes = linhas(wb["Nodes"])
    edges = linhas(wb["Edges"])
    perfis = linhas(wb["Perfis"])

    ids_nodes = {txt(r[0]) for _, r in nodes}
    ids_perfis = {txt(r[0]) for _, r in perfis}


    # ---------- V1: duas fontes por perfil (R6) ----------
    sem_fonte_ok = 0
    for _, r in perfis:
        rid, estagio, fontes = txt(r[0]), txt(r[7]), r[10]
        n = conta_fontes(fontes)
        if estagio in ESTAGIOS_SEM_PESQUISA:
            sem_fonte_ok += 1
            continue
        if n < MIN_FONTES_PERFIL:
            erro("V1", rid, f"{n} fonte(s) no perfil; o mínimo é {MIN_FONTES_PERFIL}")

    # ---------- V2: foto precisa de crédito, link e arquivo (R2) ----------
    fotos_usadas = set()
    for _, r in perfis:
        rid, arquivo, credito, link = txt(r[0]), txt(r[13]), txt(r[14]), txt(r[15])
        if not arquivo:
            continue
        fotos_usadas.add(arquivo)
        if not credito:
            erro("V2", rid, f"foto '{arquivo}' sem crédito")
        if not link:
            erro("V2", rid, f"foto '{arquivo}' sem link da fonte")
        if not (PHOTOS_DIR / arquivo).exists():
            erro("V2", rid, f"foto '{arquivo}' não existe em photos/")

    # ---------- V3: status da ligação ----------
    # O build.py trata status desconhecido como "Reportado". Um erro de
    # digitação aqui promove uma alegação a fato: por isso é ERRO, não aviso.
    for i, r in edges:
        st = txt(r[3])
        if st not in STATUS_VALIDOS:
            erro("V3", f"Edges linha {i}", f"status '{st}' não é um valor válido")

    # ---------- V4: ligação apontando para id inexistente ----------
    for i, r in edges:
        for ponta, valor in (("Source", txt(r[0])), ("Target", txt(r[1]))):
            if valor not in ids_nodes:
                erro("V4", f"Edges linha {i}", f"{ponta} '{valor}' não existe na aba Nodes")

    # ---------- V5: nó sem perfil e perfil sem nó ----------
    for nid in sorted(ids_nodes - ids_perfis):
        aviso("V5", nid, "nó sem perfil na aba Perfis")
    for pid in sorted(ids_perfis - ids_nodes):
        erro("V5", pid, "perfil sem nó correspondente na aba Nodes")

    # ---------- V6 e V7: fonte por ligação ----------
    tem_coluna_fonte = len(edges[0][1]) > 6 if edges else False
    sem_fonte = com_uma = 0
    alegadas = alegadas_ok = 0
    faltantes = []
    if tem_coluna_fonte:
        for i, r in edges:
            n = conta_fontes(r[6] if len(r) > 6 else None)
            if n == 0:
                sem_fonte += 1
            elif n == 1:
                com_uma += 1
            if txt(r[3]) == STATUS_EXIGE_FONTE_PROPRIA:
                alegadas += 1
                if n >= MIN_FONTES_LIGACAO:
                    alegadas_ok += 1
                else:
                    faltantes.append((i, n))
        # Uma linha por ligação faltante inundaria o relatório enquanto o
        # preenchimento não termina. Enquanto a regra é aviso, sai agregado;
        # quando vira erro, cada uma precisa aparecer para ser corrigida.
        if faltantes and EXIGIR_FONTE_NA_LIGACAO:
            for i, n in faltantes:
                erro("V7", f"Edges linha {i}",
                     f"ligação '{STATUS_EXIGE_FONTE_PROPRIA}' com {n} fonte(s); "
                     f"o mínimo é {MIN_FONTES_LIGACAO} (F7)")
        elif faltantes:
            amostra = ", ".join(str(i) for i, _ in faltantes[:8])
            resto = f" e mais {len(faltantes) - 8}" if len(faltantes) > 8 else ""
            aviso("V7", "Edges",
                  f"{len(faltantes)} de {alegadas} ligações "
                  f"'{STATUS_EXIGE_FONTE_PROPRIA}' abaixo de {MIN_FONTES_LIGACAO} fontes "
                  f"(F7) — linhas {amostra}{resto}")
        if sem_fonte:
            aviso("V6", "Edges",
                  f"{sem_fonte} de {len(edges)} ligações sem fonte registrada")
    else:
        aviso("V6", "Edges", "a aba não tem coluna 'Fontes'")

    # ---------- V8: foto no repositório que ninguém usa ----------
    if PHOTOS_DIR.exists():
        no_disco = {p.name for p in PHOTOS_DIR.iterdir()
                    if p.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp"}}
        for orfa in sorted(no_disco - fotos_usadas):
            aviso("V8", orfa, "arquivo em photos/ que nenhum perfil usa")

    # ---------- V9: fórmulas da aba Resumo cobrindo menos linhas que Perfis ----------
    wbf = openpyxl.load_workbook(XLSX_PATH)
    for c in wbf["Resumo"]["B"]:
        if isinstance(c.value, str) and "Perfis!" in c.value:
            m = re.search(r"Perfis!\$[A-Z]+\$\d+:\$[A-Z]+\$(\d+)", c.value)
            if m and int(m.group(1)) < len(perfis) + 1:
                aviso("V9", f"Resumo!{c.coordinate}",
                      f"fórmula cobre até a linha {m.group(1)}, mas Perfis tem "
                      f"{len(perfis) + 1}; os totais estão subcontando")
            break

    # ---------- relatório ----------
    print("=" * 62)
    print("COBERTURA")
    print("=" * 62)
    print(f"  Nós: {len(nodes)}   Ligações: {len(edges)}   Perfis: {len(perfis)}")
    st = Counter(txt(r[3]) for _, r in edges)
    for k, v in st.most_common():
        print(f"  Status '{k}': {v}")
    com_foto = sum(1 for _, r in perfis if txt(r[13]))
    pct = 100 * com_foto / len(perfis) if perfis else 0
    print(f"  Perfis com foto: {com_foto}/{len(perfis)} ({pct:.0f}%)")
    ok = sum(1 for _, r in perfis
             if conta_fontes(r[10]) >= MIN_FONTES_PERFIL
             or txt(r[7]) in ESTAGIOS_SEM_PESQUISA)
    print(f"  Perfis dentro da regra de fontes: {ok}/{len(perfis)} "
          f"({100 * ok / len(perfis):.0f}%)")
    if sem_fonte_ok:
        print(f"    (inclui {sem_fonte_ok} no estágio 'Não pesquisado', "
              f"isentos por decisão editorial)")
    if tem_coluna_fonte:
        datadas = len(edges) - sem_fonte
        print(f"  Ligações com fonte: {datadas}/{len(edges)} "
              f"({100 * datadas / len(edges):.0f}%)")
        if com_uma:
            print(f"    ({com_uma} com uma fonte só)")
        if alegadas:
            print(f"  '{STATUS_EXIGE_FONTE_PROPRIA}' com {MIN_FONTES_LIGACAO}+ fontes "
                  f"próprias (F7): {alegadas_ok}/{alegadas} "
                  f"({100 * alegadas_ok / alegadas:.0f}%)")
            if not EXIGIR_FONTE_NA_LIGACAO:
                print("    (regra ainda em aviso; vire EXIGIR_FONTE_NA_LIGACAO "
                      "quando chegar a 100%)")

    print()
    print("=" * 62)
    print(f"RESULTADO: {len(erros)} erro(s), {len(avisos)} aviso(s)")
    print("=" * 62)
    for linha in erros:
        print(linha)
    for linha in avisos:
        print(linha)

    if erros:
        print()
        print("Erro derruba o build. Aviso é a fila de revisão humana.")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
