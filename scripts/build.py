#!/usr/bin/env python3
"""Regenera o bloco <script id="graph-data"> do index.html a partir de data/mapa.xlsx.

Uso:
    python3 scripts/build.py

A planilha é a fonte de verdade: abas Nodes, Edges e Perfis. O conteúdo de
"notes" e "stages" do index.html é preservado (muda raramente; edite o HTML
direto se precisar mexer neles).
"""
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta a biblioteca openpyxl. Instale com: pip3 install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "data" / "mapa.xlsx"
HTML_PATH = ROOT / "index.html"

STATUS_MAP = {"Documentado": "doc", "Alegado ou contestado": "alleg"}
BLANK_VALUES = {"—", "-", "", None}
MIL_BLANK_VALUES = BLANK_VALUES | {"Nada nas fontes."}


def clean(v, blanks=BLANK_VALUES):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in blanks else s


def parse_sources(raw):
    """'Título URL | Título2 URL2' -> [{"t": "Título", "u": "URL"}, ...]"""
    raw = clean(raw)
    if not raw:
        return []
    out = []
    for chunk in raw.split("|"):
        chunk = chunk.strip()
        if not chunk:
            continue
        m = re.search(r"(https?://\S+)", chunk)
        if not m:
            continue
        url = m.group(1)
        title = chunk[: m.start()].strip()
        out.append({"t": title, "u": url})
    return out


def main():
    if not XLSX_PATH.exists():
        sys.exit(f"Não achei {XLSX_PATH}")
    if not HTML_PATH.exists():
        sys.exit(f"Não achei {HTML_PATH}")

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

    # ---- Nodes ----
    nodes = []
    groups_order = []
    ws = wb["Nodes"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        nid, label, typ, group, desc, _degree, *rest = row
        foto, credito, credito_url, nota = (list(rest) + [None] * 4)[:4]
        if group and group not in groups_order:
            groups_order.append(group)
        node = {
            "id": str(nid).strip(),
            "label": clean(label),
            "type": clean(typ),
            "group": clean(group),
            "desc": clean(desc),
            "degree": 0,  # filled in below from edges
        }
        if clean(foto):
            node["img"] = f"photos/{clean(foto)}"
        if clean(credito):
            node["credit"] = clean(credito)
        if clean(credito_url):
            node["creditUrl"] = clean(credito_url)
        if clean(nota):
            node["imgNote"] = clean(nota)
        nodes.append(node)

    # ---- Edges ----
    edges = []
    degree_count = {}
    ws = wb["Edges"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        source, target, label, status = row[:4]
        source, target = str(source).strip(), str(target).strip()
        edges.append({
            "source": source,
            "target": target,
            "label": clean(label),
            "status": STATUS_MAP.get(clean(status), "doc"),
        })
        degree_count[source] = degree_count.get(source, 0) + 1
        degree_count[target] = degree_count.get(target, 0) + 1

    # degree is derived from edges, not read from the sheet's formula cell
    # (cached formula values aren't reliable across spreadsheet tools)
    for n in nodes:
        n["degree"] = degree_count.get(n["id"], 0)

    # ---- Perfis ----
    profiles = {}
    stages_seen = []
    ws = wb["Perfis"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        (rid, _nome, _grupo, quem_e, onde, _vivo, situacao, estagio, milicia,
         ressalvas, fontes, checado, mudou, *resto) = row
        foto, credito, credito_url = (resto + [None, None, None])[:3]
        stage = clean(estagio)
        if stage and stage not in stages_seen:
            stages_seen.append(stage)
        p = {
            "stage": stage,
            "bio": clean(quem_e),
            "legal": clean(situacao),
            "mil": clean(milicia, MIL_BLANK_VALUES),
            "caveat": clean(ressalvas),
            "src": parse_sources(fontes),
            "date": clean(checado),
            "chg": clean(mudou),
            "loc": clean(onde),
        }
        if clean(foto):
            p["img"] = f"photos/{clean(foto)}"
        if clean(credito):
            p["credit"] = clean(credito)
        if clean(credito_url):
            p["creditUrl"] = clean(credito_url)
        profiles[str(rid).strip()] = p

    # ---- preserve notes + stages order from current HTML ----
    html = HTML_PATH.read_text(encoding="utf-8")
    m = re.search(r'<script id="graph-data"[^>]*>', html)
    if not m:
        sys.exit("Não achei o bloco <script id=\"graph-data\"> no index.html")
    start = m.end()
    end = html.find("</script>", start)
    current = json.loads(html[start:end])

    stages = current.get("stages", [])
    for s in stages_seen:
        if s not in stages:
            stages.append(s)

    data = {
        "groups": groups_order,
        "nodes": nodes,
        "edges": edges,
        "profiles": profiles,
        "stages": stages,
        "notes": current.get("notes", {}),
    }

    new_raw = json.dumps(data, ensure_ascii=False, separators=(", ", ": "))
    new_html = html[:start] + new_raw + html[end:]
    HTML_PATH.write_text(new_html, encoding="utf-8")

    old_ids = {n["id"] for n in current.get("nodes", [])}
    new_ids = {n["id"] for n in nodes}
    added = new_ids - old_ids
    removed = old_ids - new_ids

    print(f"OK: {len(nodes)} nós, {len(edges)} conexões, {len(profiles)} perfis.")
    if added:
        print(f"Novos nós: {sorted(added)}")
    if removed:
        print(f"Nós removidos: {sorted(removed)}")
    if not added and not removed:
        print("Nenhum nó adicionado/removido (só conteúdo pode ter mudado).")


if __name__ == "__main__":
    main()
